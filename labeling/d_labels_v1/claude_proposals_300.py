"""Claude's D(t) proposals for the 300 tools of the 400-tool sample that are
NOT in the 100-tool blind evaluation set (d_eval_100_ids.txt). Model output
made in chat from descriptions only, with the frozen rules in README.md
(B1, S1, A1, A2, N1, N2 v2, H1, E1, W1, W2, RD, CD, NS). Not a logged API run,
so it is a draft for human review, not the model condition of the evaluation.

Run: python claude_proposals_300.py  ->  d_rest_300_claude_proposals.xlsx
"""
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "scripts"))
from make_d_labeling import COLUMNS, EFFECTS  # noqa: E402

HERE = Path(__file__).resolve().parent
A, D, U = "affirmed", "denied", "uncertain"


def c(effect, polarity, scope, evidence, actor="tool", condition="", mode="", world="", review="no", notes=""):
    return dict(effect=effect, polarity=polarity, scope=scope, actor=actor, condition=condition, mode=mode,
                world=world, evidence=evidence, needs_review=review, notes=notes)


P = {}

# ---- batch 1 (1-60) ----
P.update({
    "sample_400:1": [
        c("write", A, "other", "Ask the user for a temporary cross-site window (30 or 60 minutes)", mode="additive",
          review="yes", notes="grants a temporary permission window; scripts run later by the agent"),
        c("delete", A, "other", "Auto-expires", actor="provider", notes="A1: the window expires by itself"),
    ],
    "sample_400:2": [
        c("read", A, "database", "Report dataset provenance, refresh time, cache origin, and record counts."),
    ],
    "sample_400:5": [
        c("write", A, "database", "Restore a previously archived product (sets it back to active).", mode="modifying"),
        c("write", D, "remote_api", "It is not re-synced to any sales channel automatically",
          notes="S1: no sync to sales channels"),
    ],
    "sample_400:6": [
        c("read", A, "remote_api", "Sequentially download all ready results for one job"),
        c("write", A, "filesystem", "download all ready results for one job into CASDA_DOWNLOAD_DIR", mode="unknown"),
        c("network", A, "remote_api", "Sequentially download all ready results for one job", world="closed",
          notes="N2 v2: an explicit download; the user's own job results"),
    ],
    "sample_400:7": [
        c("read", A, "database", "List backlog tasks and each member's current story points"),
    ],
    "sample_400:9": [
        c("write", A, "filesystem", "Replaces the entire content of the frame.", mode="overwrite"),
    ],
    "sample_400:10": [
        c("execute", A, "remote_api", "Runs a Studio draft attack on SafeBreach simulators.", review="yes",
          notes="runs an attack simulation on remote simulators"),
        c("write", A, "remote_api", "Running on all\nconnected simulators creates excessive noise on the platform",
          mode="additive", condition="all_connected=True", review="yes", notes="creates a test execution"),
    ],
    "sample_400:11": [
        c("read", U, "other", "Run one fixed, authorized operation declared with risk read.",
          notes="the operation is declared read-risk; what it reads is not said"),
    ],
    "sample_400:13": [
        c("read", A, "other", "Report vault configuration status for the Obsidian bridge"),
    ],
    "sample_400:15": [
        c("write", A, "filesystem", "Create a Markdown note with UUID inside agent.", mode="additive",
          notes="'never overwrites'"),
        c("write", A, "filesystem", "Creates parent folders", mode="additive"),
    ],
    "sample_400:16": [
        c("read", A, "database", "Browse, plan, run, and steer multi-step pipelines"),
        c("execute", A, "other", "run a stored plan by its id", review="yes", notes="runs a pipeline"),
        c("write", A, "database", "approve or deny a step, cancel, or retry a step", mode="modifying"),
    ],
    "sample_400:17": [
        c("network", A, "remote_api", "Connects to each server, retrieves tool definitions", world="unknown",
          notes="N2 v2: an explicit connection"),
        c("read", A, "remote_api", "retrieves tool definitions"),
    ],
    "sample_400:19": [
        c("write", A, "database", "create user in database", mode="additive"),
    ],
    "sample_400:20": [
        c("write", A, "messages", "Say a message in chat on every server, or on the ones listed, as the admins.",
          mode="additive"),
        c("network", U, "messages", "Say a message in chat on every server", world="unknown",
          notes="sending to game servers implies a connection but none is described"),
    ],
    "sample_400:21": [
        c("read", A, "user_data", "Get exact Safari bookmark or Reading List detail by opaque handle"),
        c("write", D, "global", "capped and read-only", notes="B1"),
        c("delete", D, "global", "capped and read-only", notes="B1"),
    ],
    "sample_400:22": [
        c("read", A, "database", "Audit the ledger's internal consistency"),
        c("write", D, "global", "Read-only.", notes="B1"),
        c("delete", D, "global", "Read-only.", notes="B1"),
    ],
    "sample_400:23": [],
    "sample_400:27": [
        c("write", A, "remote_api", "Sets a permission overwrite on a single channel for a role or member", mode="modifying"),
        c("network", U, "remote_api", "(PUT /channels/{channel_id}/permissions/{overwrite_id})", world="unknown",
          notes="N2 v2: bare method + route"),
    ],
    "sample_400:28": [
        c("write", A, "remote_api", "Close a specified size on an open position via Capital.com.", mode="modifying"),
        c("write", A, "remote_api", "the client falls back to close + reopen with the remaining size",
          condition="Capital rejects the partial DELETE", mode="modifying"),
        c("network", U, "remote_api", "If Capital rejects the partial DELETE", world="unknown",
          notes="names an API response from Capital.com; no call described as such"),
    ],
    "sample_400:30": [
        c("write", A, "database", "Update properties for an existing category or category group.", mode="modifying"),
        c("write", A, "database", "supplying children replaces the group's full child list", mode="overwrite",
          condition="children supplied for a category group"),
    ],
    "sample_400:31": [
        c("read", A, "user_data", "Score bills against your saved topics"),
        c("execute", A, "code", "using an AI model that runs on this computer", review="yes",
          notes="runs a local model"),
        c("network", D, "user_data", "Your topics and the bill text never leave the machine", notes="S1"),
    ],
    "sample_400:33": [
        c("read", A, "database", "Summarize closed local paper trades for a rolling day or week."),
        c("network", D, "remote_api", "without contacting T-Invest", notes="S1: T-Invest only"),
    ],
    "sample_400:34": [
        c("execute", U, "other", "IDB (fb-idb) must be installed", notes="needs an external program; running it is not said"),
    ],
    "sample_400:35": [
        c("read", A, "remote_api", "Returns a full account snapshot"),
    ],
    "sample_400:36": [
        c("read", A, "user_data", "Search local Apple Podcasts show metadata"),
    ],
    "sample_400:37": [
        c("read", A, "user_data", "Inspect the local computer before acting."),
    ],
    "sample_400:38": [
        c("read", A, "filesystem", "Reads local files automatically"),
    ],
    "sample_400:39": [
        c("read", A, "remote_api", "Get detailed app wishlist reporting for one date"),
    ],
    "sample_400:40": [
        c("read", A, "filesystem", "Retrieve the actual content of a specific numbered chunk from a diff file."),
    ],
    "sample_400:42": [
        c("read", U, "remote_api", "Find verified internal-link review candidates for a target URL",
          notes="RD: retrieval from the site is implied, not stated"),
    ],
    "sample_400:43": [
        c("read", A, "database", "Read the current or final status for a memory ingest run."),
    ],
    "sample_400:45": [
        c("write", A, "remote_api", "Rollback to a specific deployment", mode="modifying"),
    ],
    "sample_400:47": [
        c("write", A, "database",
          "Update a database title, description, icon, cover, inline display, parent, lock state, or trash state.",
          mode="modifying"),
        c("delete", U, "database", "trash state", review="yes", notes="moving to trash may be a delete"),
    ],
    "sample_400:48": [
        c("read", A, "remote_api", "Download a complete authorized media item locally"),
        c("write", A, "filesystem", "This writes a file", mode="additive"),
        c("network", A, "remote_api", "can use significant bandwidth", world="unknown",
          notes="an explicit download using bandwidth"),
    ],
    "sample_400:49": [
        c("read", A, "credentials", "Lists the labels of all passwords stored in the user's vault."),
        c("read", D, "other", "Passwords themselves are never returned.", notes="S1: password values"),
    ],
    "sample_400:50": [
        c("write", A, "remote_api", "Begin a browser-based connection", mode="additive", review="yes",
          notes="starts an authorization; finished by the user in a browser"),
        c("read", A, "remote_api", "You must keep polling check_connection_status", actor="other_tool"),
    ],
    "sample_400:51": [
        c("write", A, "database", "idempotently import a provider adapter's bounded CI handoff bundle", mode="additive"),
    ],
    "sample_400:52": [
        c("read", A, "filesystem", "Render one page to a PNG image."),
        c("write", D, "global", "Read-only.", notes="B1"),
        c("delete", D, "global", "Read-only.", notes="B1"),
    ],
    "sample_400:53": [
        c("read", A, "other", "Take a screenshot of the iOS Simulator."),
    ],
    "sample_400:54": [
        c("read", A, "filesystem", "Read a permitted document, heading or line range with numbered lines"),
    ],
    "sample_400:55": [
        c("read", A, "database", "Walk the tenant's entire hash chain."),
    ],
    "sample_400:57": [],
    "sample_400:58": [
        c("write", A, "other", "设置指定资源（数据库集合、云函数、存储桶）的安全规则", mode="modifying",
          notes="sets security rules on database collections, functions, buckets"),
    ],
    "sample_400:60": [
        c("read", U, "remote_api", "Check if a newer version of Syncthing is available.",
          notes="RD: where it checks is not said"),
    ],
    "sample_400:61": [
        c("delete", A, "database", "Delete a todo task."),
    ],
    "sample_400:64": [
        c("read", A, "remote_api", "Export a Part Studio as a millimeter STL"),
        c("write", U, "filesystem", "Export a Part Studio as a millimeter STL", mode="unknown",
          notes="an export may write a file or return the data"),
    ],
    "sample_400:66": [
        c("delete", A, "user_data", "Unfollow one or more artists or users.", notes="removes a follow (membership)"),
    ],
    "sample_400:67": [
        c("write", A, "filesystem", "Writes the full content of a text file, creating it (and missing folders) if needed.",
          mode="overwrite"),
        c("write", A, "filesystem", "Existing docs are updated in place via a minimal live diff", mode="modifying"),
    ],
    "sample_400:68": [
        c("write", A, "filesystem", "Replace body, header, or footer text in the active Word document in place.",
          mode="overwrite"),
    ],
    "sample_400:69": [
        c("read", A, "user_data", "Возвращает личную запись пользователя за указанную дату."),
    ],
    "sample_400:70": [
        c("read", A, "user_data", "Read the logged-in RENPHO Health account profile."),
        c("read", D, "credentials", "Credentials and tokens are excluded.", notes="S1"),
    ],
    "sample_400:72": [],
    "sample_400:73": [],
    "sample_400:74": [
        c("read", A, "other", "Check if a CAPTCHA is present on the current page."),
    ],
    "sample_400:75": [
        c("write", A, "database", "Cria conta manual.", mode="additive", condition="confirm=true"),
    ],
    "sample_400:76": [],
    "sample_400:77": [
        c("read", A, "remote_api", "Get detailed information about a podcast show."),
    ],
    "sample_400:79": [
        c("read", A, "filesystem", "Run read-only ripgrep inside the current workspace"),
        c("execute", A, "code", "Run read-only ripgrep", notes="a fixed program"),
        c("write", D, "filesystem", "read-only ripgrep", notes="S1: the search itself"),
    ],
    "sample_400:81": [
        c("write", A, "other", "Reboot a camera.", mode="modifying",
          notes="footage during the reboot is lost: data never recorded, not deleted"),
    ],
    "sample_400:82": [
        c("write", A, "other", "Upload an image (event cover, org logo, avatar) to the media bucket.", mode="additive"),
        c("network", U, "other", "Upload an image (event cover, org logo, avatar) to the media bucket.", world="closed",
          notes="an upload to the server's own bucket; transport not described"),
    ],
})

# ---- batch 2 (61-120) ----
P.update({
    "sample_400:83": [c("read", A, "remote_api", "Execute a saved TelemetryDeck insight and return results.",
                        notes="runs a saved query, not code")],
    "sample_400:85": [c("write", A, "other", "Turn a camera's smart detections on or off", mode="modifying")],
    "sample_400:87": [],
    "sample_400:88": [c("read", A, "database", "Registration, team, and check-in counts for a single event.")],
    "sample_400:89": [
        c("execute", A, "code", "The dash's confirmed reap (dash-reap.sh --yes)", notes="a fixed script"),
        c("delete", A, "filesystem", "remove the worktree when clean", condition="worktree clean"),
        c("write", A, "other", "close the window", mode="modifying"),
    ],
    "sample_400:90": [
        c("execute", A, "remote_api", "Runs the advisor's fixed read-only shell probe on an SSM-managed Linux instance",
          notes="a fixed program on a remote instance"),
        c("read", A, "remote_api", "returns memory, disks, load, top processes, containers"),
        c("write", D, "remote_api", "fixed read-only shell probe", notes="S1: the probe"),
        c("network", A, "remote_api", "through Run Command", world="closed",
          notes="an explicit remote call to the user's instance"),
    ],
    "sample_400:91": [
        c("read", A, "filesystem", "solo procesa los archivos"),
        c("write", A, "database", "Importa un export de Apple Salud / Health Auto Export a una base local (SQLite).",
          mode="additive"),
        c("delete", A, "database", "para borrar y reconstruir desde cero", condition="full_rebuild=True"),
    ],
    "sample_400:92": [c("read", A, "database", "List all taxonomy definitions (e.g. categories, tags).")],
    "sample_400:93": [
        c("read", A, "remote_api", "Fetch the still image Protect captured for an event"),
        c("network", A, "remote_api", "Fetch the still image Protect captured for an event", world="closed",
          notes="an explicit fetch from the user's Protect console"),
        c("write", A, "filesystem", "Writes it to disk and returns the path by default", mode="additive",
          condition="default output"),
    ],
    "sample_400:94": [
        c("network", A, "remote_api", "Fetch external Swift-DocC documentation by absolute https URL", world="open",
          notes="N1: any URL"),
        c("read", A, "remote_api", "Fetch external Swift-DocC documentation by absolute https URL"),
    ],
    "sample_400:95": [c("read", A, "other",
                        "Report the wifi radio state, the current network, and scan for available ones.")],
    "sample_400:96": [
        c("write", A, "remote_api", "Cancel an outstanding Bill Manager invoice by its external reference.",
          mode="modifying", notes="irreversible: 'cannot be restored through this tool'"),
        c("network", U, "remote_api", "Returns the Daraja response", world="unknown", notes="names an API response"),
    ],
    "sample_400:97": [c("read", A, "remote_api", "Returns Groupon deals and low-price options sorted cheapest first.")],
    "sample_400:98": [c("read", A, "remote_api", "uMap API: read IRE rules for a kind")],
    "sample_400:99": [c("write", A, "database", "Mark a todo as complete or incomplete.", mode="modifying")],
    "sample_400:101": [c("read", A, "database",
                         "Read compact browser evidence records or an investigation evidence graph.")],
    "sample_400:102": [c("write", A, "filesystem",
                         "Generate a new Activity Kotlin file from a template and add it to an existing Android project",
                         mode="additive")],
    "sample_400:103": [
        c("read", A, "remote_api", "Query GBIF for georeferenced threatened & endangered species occurrences."),
        c("network", A, "remote_api", "Query GBIF", world="unknown", notes="N2 v2: an explicit query to GBIF"),
    ],
    "sample_400:104": [c("read", A, "other", "Returns visible text from the whole page or from a selected element.")],
    "sample_400:106": [
        c("network", A, "remote_api", "via online provider lookup", world="unknown"),
        c("read", A, "remote_api", "on Open Library and Google Books"),
        c("write", D, "global", "never writes", notes="'Preview only'"),
    ],
    "sample_400:108": [c("write", A, "remote_api", "Create a new sprint on a board.", mode="additive")],
    "sample_400:109": [c("read", U, "database", "to archive objects", notes="resolves an id to archived objects")],
    "sample_400:110": [c("write", A, "other", "Bind one scanned dimension annotation to likely measured geometry.",
                         mode="additive")],
    "sample_400:111": [c("read", A, "database", "fetch full episode by ID")],
    "sample_400:112": [],
    "sample_400:114": [
        c("read", A, "remote_api",
          "Verify a Safaricom number: registration date, SIM age, last swap date and a hashed IMSI"),
        c("write", D, "global", "Read-only.", notes="B1"),
        c("delete", D, "global", "Read-only.", notes="B1"),
        c("network", U, "remote_api", "as returned by the API", world="unknown"),
    ],
    "sample_400:115": [c("read", A, "filesystem", "List paragraphs with their styles in the active Word document")],
    "sample_400:116": [c("read", A, "remote_api",
                         "Display details about the metadata types that are enabled for your org.")],
    "sample_400:117": [],  # RD: searches bundled (built-in) threat rules
    "sample_400:118": [
        c("write", A, "messages", "Send an email directly, restricted to internal recipients only (same domain).",
          mode="additive"),
        c("network", A, "messages", "Send an email directly", world="closed", notes="internal recipients only"),
    ],
    "sample_400:119": [],
    "sample_400:120": [c("write", A, "other", "Auto-duplicates — the original is preserved", mode="additive",
                         notes="paints a duplicated layer")],
    "sample_400:121": [
        c("write", A, "messages", "Apply an approved exact-chat Messages text or file send", mode="additive"),
        c("read", A, "messages", "verify local read-back"),
    ],
    "sample_400:123": [
        c("write", A, "other", "Complete one authorized begin_web_search lease.", mode="modifying"),
        c("read", A, "remote_api", "Server reads each exact page against original requirements."),
        c("network", A, "remote_api", "Server reads each exact page", world="open",
          notes="pages at submitted merchant URLs"),
        c("network", D, "user_data", "Never upload the user's reference image.", notes="S1"),
    ],
    "sample_400:124": [
        c("read", A, "remote_api", "Reports every attack and simulator the console measured at exactly ZERO"),
        c("write", D, "global", "and changes nothing"),
        c("delete", D, "global", "and changes nothing"),
        c("execute", D, "remote_api", "WITHOUT running it", notes="S1: the scenario"),
    ],
    "sample_400:126": [],  # scans the caller's input; the per-call charge is not modelled as a tool effect
    "sample_400:127": [c("read", A, "remote_api", "Discover all searchable fields for any GLPI item type.")],
    "sample_400:130": [
        c("read", A, "remote_api", "Retrieve historical WHOIS or BGP data for a single indicator."),
        c("network", U, "remote_api", "On upstream failure", world="unknown", notes="names an upstream source"),
    ],
    "sample_400:133": [c("write", A, "filesystem", "Find and replace text in the document.", mode="modifying")],
    "sample_400:134": [c("read", A, "remote_api", "Get all issues for a specific project.")],
    "sample_400:135": [
        c("delete", A, "database", "Permanently delete a term from a taxonomy."),
        c("delete", A, "database", "Any content tagged with this term loses the association."),
    ],
    "sample_400:136": [
        c("read", A, "other", "validates the device's saved configuration on the dashboard"),
        c("write", D, "global", "without modifying anything"),
        c("delete", D, "global", "without modifying anything"),
        c("execute", U, "code", "``esphome`` to be installed", condition="local file path given",
          review="yes", notes="runs the ESPHome validator?"),
    ],
    "sample_400:137": [c("read", A, "database",
                         "Read a cached official page by document ID or approved documentation URL.")],
    "sample_400:140": [
        c("read", A, "remote_api", "Get the hierarchical table of contents for a CFR title or subset."),
        c("network", U, "remote_api", "(returns 400)", world="unknown", notes="an HTTP status hints at an API"),
    ],
    "sample_400:141": [
        c("read", A, "remote_api", "Use when you want a US mutual fund's profile"),
        c("network", U, "remote_api", "GET /funds/snapshot/{bw_fund_id}", world="unknown",
          notes="N2 v2: method + route"),
        c("write", A, "other", "Price: $0.05 for the snapshot", actor="provider", mode="additive", notes="billing"),
    ],
    "sample_400:142": [c("read", A, "user_data", "segun el perfil y el entreno")],
    "sample_400:144": [],
    "sample_400:146": [
        c("read", A, "remote_api", "Synchronize a screen from Google Stitch to local HTML files."),
        c("write", A, "filesystem", "Synchronize a screen from Google Stitch to local HTML files.", mode="unknown"),
        c("network", A, "remote_api", "Synchronize a screen from Google Stitch", world="closed",
          notes="an explicit transfer from the user's project"),
    ],
    "sample_400:147": [c("delete", A, "database", "Delete a saved workflow by ID.")],
    "sample_400:150": [],  # RD: returns built-in documentation
    "sample_400:151": [
        c("read", A, "database", "Read bounded shopping references and historical observations for this host task"),
        c("read", D, "other", "This tool cannot access another task", notes="S1"),
    ],
    "sample_400:152": [],
    "sample_400:153": [
        c("read", A, "remote_api", "Queries the shared body of debug knowledge"),
        c("network", U, "remote_api", "Recall known solutions from the public debug-knowledge commons.",
          world="unknown"),
        c("write", A, "remote_api", "use 'remember' to register the problem", actor="other_tool", mode="additive"),
    ],
    "sample_400:155": [c("read", A, "database", "Lists projects the signed-in member can see.")],
    "sample_400:156": [
        c("write", A, "other", "create/delete/ungroup/add_layer/set_blend_mode", mode="modifying"),
        c("delete", A, "other", "delete removes the group AND everything inside it",
          condition="op=delete, confirm:true"),
    ],
    "sample_400:157": [c("read", A, "remote_api", "List and filter Odoo models.")],
    "sample_400:158": [
        c("read", A, "filesystem", "DAEDALUS takes a sample path from"),
        c("execute", A, "code", "executes it while capturing filesystem",
          notes="runs the caller's sample in a VM"),
        c("write", A, "other", "a VM is created and", mode="additive"),
        c("delete", A, "other", "then destroys the VM", notes="its own VM (cleanup)"),
        c("write", A, "filesystem", "pre/post filesystem tars and the report in the store", mode="additive"),
    ],
    "sample_400:160": [
        c("delete", A, "database", "Soft-deletes one to 100 missions atomically"),
        c("delete", A, "database", "Deleting a mission also soft-deletes its live objectives."),
    ],
    "sample_400:161": [c("read", A, "filesystem", "Fast fuzzy file search with frecency ranking.")],
    "sample_400:162": [
        c("write", A, "other", "Write contract.", mode="unknown", review="yes", notes="a contract write; target not said"),
        c("write", A, "other", "Costs $0.01 USDC per call.", actor="provider", mode="additive", notes="billing"),
    ],
})

# ---- batch 3 (121-180) ----
P.update({
    "sample_400:163": [c("read", A, "user_data",
                         "List selected local Calendar target event metadata in an explicit date window.")],
    "sample_400:164": [c("read", A, "database", "Fulfillment providers and sales channels this account may connect")],
    "sample_400:168": [
        c("read", A, "remote_api", "Get the text content of a doc, presentation, whiteboard"),
        c("write", D, "global", "without making changes"),
        c("delete", D, "global", "without making changes"),
    ],
    "sample_400:172": [c("write", A, "database",
                         "Update frame metadata: rename it, move/resize it, or move it to another page using pageId.",
                         mode="modifying")],
    "sample_400:174": [
        c("network", A, "remote_api", "Fetch the official Anthropic Claude Code docs (llms-full.txt)", world="unknown",
          notes="an explicit fetch from a fixed source"),
        c("read", A, "remote_api", "Fetch the official Anthropic Claude Code docs (llms-full.txt)"),
        c("write", A, "filesystem", "overwrites previous baseline AND current", mode="overwrite",
          notes="its own local snapshot"),
    ],
    "sample_400:175": [c("write", A, "database", "Create a pending local approval request for a risky action.",
                         mode="additive")],
    "sample_400:176": [
        c("read", A, "database", "Find up to three ClueXP provider companies that can serve a skill at a location."),
        c("write", D, "global", "Nothing is booked, dispatched, or confirmed by this tool"),
    ],
    "sample_400:177": [
        c("read", A, "other", "Read the text in a design with local OCR (tesseract) when available"),
        c("execute", A, "code", "local OCR (tesseract)", condition="when available", notes="a fixed program"),
    ],
    "sample_400:179": [c("write", A, "remote_api",
                         "This will save all the changes made to the design in the specified editing transaction.",
                         mode="modifying")],
    "sample_400:180": [
        c("read", A, "remote_api", "List prediction market events from Polymarket."),
        c("write", A, "other", "Costs $0.001 USDC.", actor="provider", mode="additive", notes="billing"),
    ],
    "sample_400:181": [c("read", A, "database", "Search working memory and summaries in the active session.")],
    "sample_400:182": [c("read", A, "remote_api", "Get runtime logs for a deployment")],
    "sample_400:184": [
        c("read", A, "user_data", "The local algorithm uses saved tastes and filters"),
        c("write", A, "database", "Directo recommends games and saves them in history.", mode="additive",
          condition="Directo mode"),
        c("network", U, "remote_api", "Codex consumes quota and interprets messages.", world="unknown",
          condition="Codex mode", notes="N2 v2: an operation-linked quota"),
    ],
    "sample_400:187": [c("read", A, "remote_api", "Get valid values for selection (dropdown) fields on an Odoo model.")],
    "sample_400:188": [c("read", A, "remote_api",
                         "Get a preview image of a LiveDB material rendered on a material ball.")],
    "sample_400:189": [c("read", A, "database",
                         "Search across KTX wiki pages, semantic-layer sources, measures, dimensions, raw tables, and columns.")],
    "sample_400:190": [],
    "sample_400:191": [c("write", A, "other", "Generate a short video with Veo, Kling, or Seedance 2.5", mode="additive",
                         notes="starts a generation job ('async job queue')")],
    "sample_400:193": [c("read", A, "remote_api", "Search across the pre.dev knowledge base")],
    "sample_400:194": [c("write", A, "remote_api",
                         "Create a stack (an ordered series of dependent changes) on a base branch.", mode="additive")],
    "sample_400:195": [
        c("execute", A, "code", "Run a single Cypress spec file and return results.", notes="a fixed test runner"),
        c("network", U, "other", "The dev server must be running at baseUrl (default localhost:8080)", world="closed"),
    ],
    "sample_400:197": [c("write", A, "filesystem",
                         "Moves and/or renames a file or folder inside a project, creating target folders as needed.",
                         mode="modifying")],
    "sample_400:198": [c("read", A, "other", "Read the Cue MCP/agent manifest for a session.")],
    "sample_400:199": [c("read", A, "database", "Find unblocked tasks to work on.")],
    "sample_400:200": [c("read", A, "remote_api", "Run a pre-defined common GraphQL query.",
                         notes="the listed queries all read")],
    "sample_400:201": [c("read", A, "database", "Find recorded work and its history by task, session, path, branch, or commit.")],
    "sample_400:202": [c("read", A, "remote_api", "Get the list of cryptocurrencies supported for manual tracking.")],
    "sample_400:203": [
        c("read", A, "other", "Verify RENPHO login and scale discovery"),
        c("read", D, "other", "without returning measurements.", notes="S1: measurements"),
    ],
    "sample_400:204": [c("read", A, "remote_api", "Get public information for one Steam game server account")],
    "sample_400:205": [c("read", A, "filesystem", "Show chronological timeline of frames.")],
    "sample_400:206": [c("read", A, "remote_api", "Diff a single CFR section between two dates, paragraph by paragraph.")],
    "sample_400:207": [c("write", A, "database", "Create a new content item in a collection.", mode="additive",
                         notes="created as draft")],
    "sample_400:209": [c("read", A, "other",
                         "List installed or available Android SDK packages, system images, and build tools")],
    "sample_400:210": [c("read", A, "database", "获取某个 API 接口的完整详情")],
    "sample_400:212": [c("write", A, "remote_api", "Replace a DNS load balancer's full specification (PUT).",
                         mode="overwrite")],
    "sample_400:213": [c("read", A, "other", "Remote devices that tried to connect but are not yet configured.")],
    "sample_400:214": [
        c("network", A, "remote_api", "Executes any Canvas API endpoint directly.", world="unknown"),
        c("read", A, "remote_api", "Executes any Canvas API endpoint directly."),
        c("write", A, "remote_api", "Non-GET methods CREATE, MODIFY, or DELETE real data in Canvas immediately",
          condition="non-GET method", mode="modifying"),
        c("delete", A, "remote_api", "Non-GET methods CREATE, MODIFY, or DELETE real data in Canvas immediately",
          condition="non-GET method"),
        c("network", D, "remote_api", "preview the prepared request without sending it", condition="dry_run=true",
          notes="CD"),
    ],
    "sample_400:215": [c("write", A, "remote_api", "Subscribe to notifications for a label.", mode="additive")],
    "sample_400:216": [c("write", A, "database", "Move an Utsav event through its lifecycle", mode="modifying")],
    "sample_400:217": [c("read", A, "database", "standing constraints from past decisions")],
    "sample_400:219": [
        c("read", A, "remote_api", "Query NOAA All_Species_Ranges records within the ROI."),
        c("network", A, "remote_api", "Query NOAA All_Species_Ranges records", world="unknown",
          notes="N2 v2: an explicit query"),
    ],
    "sample_400:220": [c("read", A, "user_data", "Return the signed-in AKCOE user's profile")],
    "sample_400:221": [c("write", A, "other", "Fine-tune a single SEGA dimension.", mode="modifying")],
    "sample_400:223": [c("write", A, "user_data", "Update the visible scoring methodology.", mode="modifying")],
    "sample_400:224": [
        c("write", A, "remote_api", "Update the stop loss on an open position on Capital.com.", mode="modifying"),
        c("write", A, "database", "also updates the local DB audit trail for both legs", condition="trade_id provided",
          mode="modifying"),
    ],
    "sample_400:225": [c("write", A, "other", "Repay borrowed assets to AAVE v3.", mode="modifying",
                         notes="other = on-chain debt position")],
    "sample_400:226": [c("read", U, "other", "this checks the render on the garment",
                         notes="RD: whether the render is fetched or passed in is not said")],
    "sample_400:228": [
        c("write", A, "database", "Record bounded private feedback for this credential.", mode="additive"),
        c("write", D, "remote_api", "automatic issue publishing", notes="S1: no issue publishing"),
    ],
    "sample_400:229": [c("write", A, "remote_api", "Create a new Google Doc with optional initial content.",
                         mode="additive")],
    "sample_400:230": [
        c("read", A, "remote_api", "Check every signature in an envelope at once."),
        c("write", D, "global", "Read-only, no quota.", notes="B1"),
        c("delete", D, "global", "Read-only, no quota.", notes="B1"),
    ],
    "sample_400:231": [c("read", A, "database", "Every pool the inventory knows")],
    "sample_400:232": [],
    "sample_400:233": [
        c("read", A, "remote_api", "Execute a read-only GraphQL query against the Politigraph API"),
        c("network", A, "remote_api", "against the Politigraph API", world="unknown", notes="an explicit query"),
        c("write", D, "remote_api", "read-only GraphQL query", notes="S1: the query"),
    ],
    "sample_400:234": [
        c("write", D, "global", "Read-only; it changes nothing.", notes="B1; the QR is computed"),
        c("delete", D, "global", "Read-only; it changes nothing.", notes="B1"),
    ],
    "sample_400:235": [c("read", A, "user_data", "List detected meetings with duration, app, and attendees.")],
    "sample_400:237": [c("read", A, "remote_api", "获取笔记评论")],
    "sample_400:238": [c("write", D, "filesystem", "without mutating local files.", notes="plans only")],
    "sample_400:241": [c("write", D, "other", "Valida una pipeline di workflow senza salvarla",
                         notes="checks against a bundled manifest (RD: not a read)")],
    "sample_400:242": [
        c("execute", A, "code", "Run the model's integrity audit", notes="runs under the server's own interpreter"),
        c("write", D, "global", "Read-only;", notes="B1"),
        c("delete", D, "global", "Read-only;", notes="B1"),
    ],
    "sample_400:243": [c("read", A, "database",
                         "Search 2,583 bookable Hawaii tours and activities by keyword, island, price range.")],
})

# ---- batch 4 (181-240) ----
P.update({
    "sample_400:244": [c("read", A, "filesystem", "List all Cypress e2e spec files in the project.")],
    "sample_400:247": [c("read", A, "database", "Lists tags with their kind.")],
    "sample_400:250": [c("read", U, "remote_api", "Compare key metrics across multiple stocks.",
                         notes="RD: where the metrics come from is not said")],
    "sample_400:251": [
        c("write", A, "database", "with the daemon write queue and audit flow", mode="unknown"),
        c("network", U, "remote_api", "via the /sync endpoint", world="unknown", notes="N2 v2: a bare route"),
    ],
    "sample_400:252": [c("read", A, "filesystem", "Get workbook information (name, path, sheet count) in Microsoft Excel")],
    "sample_400:253": [c("read", A, "other", "Retrieve system performance metrics")],
    "sample_400:254": [],
    "sample_400:256": [
        c("read", A, "remote_api", "Read-only listing of Kanban AI projects"),
        c("write", D, "global", "Does not modify any data."),
        c("delete", D, "global", "Does not modify any data."),
    ],
    "sample_400:258": [c("read", U, "database", "O GerminaBot tem conhecimento sobre as classes do modelo YOLO",
                         notes="answers from the bot's knowledge; retrieval not described")],
    "sample_400:259": [],  # the description is an unfilled template placeholder
    "sample_400:260": [
        c("network", A, "remote_api", "Fetches the URL", world="open", notes="N1"),
        c("read", A, "remote_api", "Fetches the URL, extracts readable content"),
        c("write", A, "database", "creates a note with metadata labels", mode="additive"),
    ],
    "sample_400:261": [c("read", A, "remote_api",
                         "Counts distinct devices whose logs contain text during an optional date range.")],
    "sample_400:263": [c("read", A, "user_data", "Get calendar events for a specific date")],
    "sample_400:265": [c("write", A, "remote_api", "修改 Apifox 中现有的接口文档", mode="modifying")],
    "sample_400:266": [
        c("read", A, "filesystem", "Validate completed replay and experiment artifacts"),
        c("write", A, "filesystem", "write a PR-ready Markdown evidence report with a SHA-256 digest", mode="additive"),
    ],
    "sample_400:268": [c("read", A, "remote_api", "List teams in an Asana workspace.")],
    "sample_400:270": [
        c("write", A, "remote_api", "Change only the category of one existing transaction.", mode="modifying"),
        c("write", D, "other", "This tool cannot change amount, date, payee", notes="S1: other fields"),
    ],
    "sample_400:271": [
        c("write", A, "remote_api", "Add a followup message to an existing GLPI ticket.", mode="additive"),
        c("write", A, "messages", "GLPI notifications will be triggered automatically.", actor="provider",
          mode="additive"),
    ],
    "sample_400:272": [
        c("read", A, "remote_api", "Read data from a Google Sheet or uploaded Excel"),
        c("write", A, "remote_api", "a temporary Google Sheets copy is created and deleted after reading",
          condition="non-native format", mode="additive"),
        c("delete", A, "remote_api", "a temporary Google Sheets copy is created and deleted after reading",
          condition="non-native format", notes="its own temporary copy"),
    ],
    "sample_400:273": [
        c("read", A, "remote_api", "Query legacy VeVote (Stakeholder and Steering Committee Governance contracts) historic proposals"),
        c("network", U, "remote_api", "/api/v1/vevote/historic-proposals", world="unknown", notes="a bare route"),
    ],
    "sample_400:275": [c("read", A, "database", "List live reviews or show one review's step progress.")],
    "sample_400:276": [
        c("read", A, "remote_api", "Use when you need the daily dividend-adjusted total returns"),
        c("network", U, "remote_api", "from GET /ticker-returns", world="unknown", notes="method + route"),
        c("write", A, "other", "Price: $0.02 per call", actor="provider", mode="additive", notes="billing"),
    ],
    "sample_400:277": [c("write", A, "other", "the user explicitly approves the reviewed cold email brief in the app",
                         mode="modifying", notes="records the approval")],
    "sample_400:278": [c("read", A, "remote_api", "Get details for a specific Railway service")],
    "sample_400:279": [c("write", A, "remote_api", "Pause the current Spotify playback.", mode="modifying")],
    "sample_400:280": [
        c("write", A, "remote_api", "Create a GitHub issue", mode="additive"),
        c("write", A, "database", "automatically create a linked journal entry documenting the issue creation",
          mode="additive"),
    ],
    "sample_400:281": [c("write", A, "filesystem", "Create a folder on this computer for tracking legislation.",
                         mode="additive")],
    "sample_400:282": [c("read", A, "database", "List all available cheesecakes on The Cheesecake Method storefront")],
    "sample_400:283": [
        c("network", A, "remote_api", "使用联网来进行信息检索", world="open", notes="web search, any URL"),
        c("read", A, "remote_api", "也可以直接输入网址获取网页内容"),
    ],
    "sample_400:284": [c("write", A, "database", "Persist a short project note", mode="additive")],
    "sample_400:285": [
        c("read", A, "filesystem", "Content/layout issues."),
        c("write", D, "global", "Read-only.", notes="B1"),
        c("delete", D, "global", "Read-only.", notes="B1"),
    ],
    "sample_400:286": [c("write", A, "database", "Process an upload staged through POST /mcp/uploads.", mode="additive",
                         notes="imports the staged activity")],
    "sample_400:288": [c("write", A, "other", "Configure digital trigger source and pattern.", mode="modifying")],
    "sample_400:289": [c("write", A, "remote_api", "Generate a new UI screen in Google Stitch from a text prompt.",
                         mode="additive")],
    "sample_400:290": [c("read", A, "filesystem", "List the model files in the Flow models folder")],
    "sample_400:292": [
        c("write", A, "database", "Ensure a goal-local Beads database and local HTML status board exist", mode="additive"),
        c("write", D, "database", "never reinitialized", condition="database already exists", notes="CD"),
    ],
    "sample_400:294": [c("read", A, "database", "Fetch a single document by id, returning text + metadata for citation.")],
    "sample_400:295": [c("write", A, "database", "Change the kind of an existing dependency edge between two services.",
                         mode="modifying")],
    "sample_400:297": [c("read", A, "remote_api", "Retrieve transaction history for a VeChain address with filtering support.")],
    "sample_400:298": [c("read", A, "other", "Retrieve container stdout or boot logs.")],
    "sample_400:300": [
        c("read", A, "remote_api", "Execute a GitLab GraphQL query."),
        c("network", A, "remote_api", "Execute a GitLab GraphQL query.", world="unknown", notes="an explicit query"),
        c("write", U, "remote_api", "Execute a GitLab GraphQL query.", mode="unknown", review="yes",
          notes="any query string: a mutation is possible"),
    ],
    "sample_400:301": [c("write", A, "filesystem", "Creates a new modified .docx file", mode="additive")],
    "sample_400:303": [
        c("read", A, "database", "Fetch and claim any open human feedback requests on a canvas."),
        c("write", A, "database", "Claiming assigns the requests to you", mode="modifying"),
    ],
    "sample_400:304": [
        c("read", A, "remote_api", "List training datasets in the caller's account."),
        c("network", U, "remote_api", "GET /prod/v1/trainers/datasets", world="unknown", notes="method + route"),
    ],
    "sample_400:305": [c("read", A, "remote_api",
                         "List live GitHub branches and their canonical head SHAs for a connected repository.")],
    "sample_400:307": [
        c("read", A, "remote_api", "解析微信公众号文章，提取完整内容"),
        c("network", U, "remote_api", "微信公众号文章 URL（必填）", world="open",
          notes="takes an article URL; fetching it is implied, not said"),
    ],
    "sample_400:309": [c("read", A, "database", "List the imported accounts")],
    "sample_400:310": [c("read", U, "remote_api", "Return an allowlisted Acquisition.gov RFO FAQ or guidance resource.",
                         notes="RD: fetched or bundled is not said")],
    "sample_400:311": [c("delete", A, "other",
                         "Clear routing on the PCB — all traces, a single net, or just unrouted connections.")],
    "sample_400:312": [c("network", D, "remote_api", "locally, without inference or an API key",
                         notes="resolves from built-in data")],
    "sample_400:313": [c("write", A, "other", "Prepare VLM-friendly artifacts and schema", mode="additive")],
    "sample_400:314": [],
    "sample_400:315": [c("read", A, "remote_api", "Returns a paginated cross-site list of downtime incidents")],
    "sample_400:316": [c("read", A, "database",
                         "Returns: title, project association, storage label, updated time, relationship count.")],
    "sample_400:317": [],
    "sample_400:318": [
        c("read", A, "database", "FTS5 search on ALL indexed content"),
        c("execute", A, "code", "Query $spec with JS.", condition="code param", notes="runs the caller's JS"),
        c("write", A, "database", "Use storeAs to save results and return summary only", condition="storeAs",
          mode="additive"),
    ],
    "sample_400:321": [
        c("write", A, "remote_api", "Adds the bot's own reaction to a message", mode="additive"),
        c("network", U, "remote_api", "(PUT /channels/{id}/messages/{id}/reactions/{emoji}/@me)", world="unknown",
          notes="method + route"),
    ],
    "sample_400:323": [c("read", A, "other",
                         "Check the active platform backend and whether the local display/native permissions are ready.")],
    "sample_400:324": [c("read", A, "remote_api", "Get game server player stats for one app and time range")],
    "sample_400:328": [
        c("read", A, "other", "inspects the exact bound window if available"),
        c("write", D, "global", "Read-only Phase 9 preflight", notes="B1"),
        c("delete", D, "global", "Read-only Phase 9 preflight", notes="B1"),
    ],
})


def tools_in_order():
    """(tool_id, server, tool, description) for the 300 tools, in sheet order."""
    ev = set((HERE / "d_eval_100_ids.txt").read_text(encoding="utf-8").split())
    ws = openpyxl.load_workbook(HERE / "d_sample_400_menna.xlsx").active
    seen = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[1] not in ev and r[1] not in seen:
            seen[r[1]] = (r[2], r[3], r[4] or "")
    return [(k, *v) for k, v in seen.items()]


def main():
    tools = tools_in_order()
    rows, bad, done = [], [], 0
    for tool_id, server, tool, desc in tools:
        if tool_id not in P:
            continue
        done += 1
        for e in EFFECTS:
            mine = [x for x in P[tool_id] if x["effect"] == e]
            for x in mine or [dict(effect=e, polarity="not_stated")]:
                if x["polarity"] != "not_stated" and x["evidence"] not in desc:
                    bad.append((tool_id, e, x["evidence"][:60]))
                rows.append([tool_id, server, tool, desc, e] + [x.get(k, "") for k in COLUMNS[6:]])
    if bad:
        raise SystemExit(f"evidence not in description: {bad}")
    unknown = set(P) - {t[0] for t in tools}
    if unknown:
        raise SystemExit(f"ids not in the 300: {sorted(unknown)}")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "claude_proposals"
    ws.append(COLUMNS)
    for i, r in enumerate(rows, 1):
        ws.append([i] + r)
    wb.save(HERE / "d_rest_300_claude_proposals.xlsx")
    print(f"{done}/{len(tools)} tools proposed, {sum(len(v) for v in P.values())} claims; all evidence found")


if __name__ == "__main__":
    main()
