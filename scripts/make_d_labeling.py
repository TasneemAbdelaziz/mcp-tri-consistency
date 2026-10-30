"""Sheets for labeling D(t), what a tool's DESCRIPTION claims, in the shared
claim-level schema (docs/Review_D_extraction_for_Menna.docx v1.1, section 3).

1. D-calibration, second annotator: the 30 tools Menna labelled
   (D:/Downloads/Menna_Reviewed_Proposals.csv), as blank rows for Tasneem,
   so the two can be compared (kappa) and the disagreements settled.
2. D-labels for the 400-tool sample of labeling/sample_v1: these tools have
   code, so the same sample gives C<->A, D<->A and C<->D.

Annotators see the server, tool name and description only: no annotations,
no code, no model output, no other annotator's labels.

Usage: python make_d_labeling.py [--menna-csv PATH]
"""
import argparse
import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "labeling/d_labels_v1"
EFFECTS = ["read", "write", "delete", "execute", "network"]

# one vocabulary for both tracks (fixed lists, so kappa and the D<->A rules can be computed)
VOCAB = {
    "polarity": ["not_stated", "affirmed", "denied", "uncertain"],
    "actor": ["tool", "provider", "other_tool", "user"],
    "scope": ["global", "filesystem", "database", "remote_api", "credentials", "messages", "code", "user_data", "other"],
    "mode": ["additive", "modifying", "overwrite", "unknown"],
    "world": ["open", "closed", "unknown"],
    "needs_review": ["no", "yes"],
}
COLUMNS = ["row", "tool_id", "server", "tool", "description", "effect", "polarity", "scope", "actor",
           "condition", "mode", "world", "evidence", "needs_review", "notes"]


def write_sheet(path, tools):
    """tools: [(tool_id, server, tool, description)] -> one row per effect;
    an annotator adds a row (copy one) for a second claim on the same effect."""
    wb = Workbook()
    ws = wb.active
    ws.title = "labels"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
    n = 0
    for tool_id, server, tool, desc in tools:
        for e in EFFECTS:
            n += 1
            ws.append([n, tool_id, server, tool, desc, e] + [""] * (len(COLUMNS) - 6))
    last = ws.max_row + 300  # room for added claim rows
    for name, options in VOCAB.items():
        col = ws.cell(row=1, column=COLUMNS.index(name) + 1).column_letter
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{last}")
    widths = {"row": 5, "tool_id": 28, "server": 24, "tool": 22, "description": 70, "effect": 9, "polarity": 12,
              "scope": 12, "actor": 11, "condition": 18, "mode": 11, "world": 9, "evidence": 40, "needs_review": 8,
              "notes": 40}
    for name, w in widths.items():
        ws.column_dimensions[ws.cell(row=1, column=COLUMNS.index(name) + 1).column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        row[COLUMNS.index("description")].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "F2"
    wb.save(path)


def menna_tools(path):
    seen = {}
    for r in csv.DictReader(open(path, encoding="utf-8-sig")):
        seen.setdefault(r["tool_id"], r["description"])
    # server / tool from census_v2 (tool_id = 1-based line number of tools_full.jsonl)
    import gzip
    want = {int(t.split(":")[1]) - 1: t for t in seen}
    names = {}
    with gzip.open(ROOT / "data/raw/census_v2/tools_full.jsonl.gz", "rt", encoding="utf-8") as f:
        for i, line in enumerate(f):
            if i in want:
                d = json.loads(line)
                if (d.get("description") or "").strip() == seen[want[i]].strip():
                    names[want[i]] = (d["server_name"], d["tool_name"])
    return [(t, *names.get(t, ("", "")), desc) for t, desc in seen.items()]


def sample_tools():
    key = list(csv.DictReader(open(ROOT / "labeling/sample_v1/sample_400_key.csv", encoding="utf-8")))
    wanted = {(r["repo"], r["tool"]) for r in key}
    full = {}
    for path in (ROOT / "data/results/scan").glob("*.jsonl"):
        for line in open(path, encoding="utf-8"):
            t = json.loads(line)
            k = (t["repo"], t["tool"])
            if k in wanted and t.get("sha"):
                d = t.get("description")
                full.setdefault(k, d if isinstance(d, str) else "")
    out, missing = [], 0
    for r in key:
        desc = full.get((r["repo"], r["tool"]), r["description"])
        missing += not desc.strip()
        out.append((f"sample_400:{r['id']}", r["repo"], r["tool"], desc or "(no description in code)"))
    return out, missing


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--menna-csv", type=Path, default=Path("D:/Downloads/Menna_Reviewed_Proposals.csv"))
    args = ap.parse_args()
    OUT.mkdir(parents=True, exist_ok=True)

    calib = menna_tools(args.menna_csv)
    write_sheet(OUT / "d_calibration_30_tasneem.xlsx", calib)
    print(f"d_calibration_30_tasneem.xlsx: {len(calib)} tools ({sum(1 for c in calib if c[1])} with server/tool names)")

    tools, missing = sample_tools()
    write_sheet(OUT / "d_sample_400_menna.xlsx", tools)
    write_sheet(OUT / "d_sample_400_tasneem.xlsx", tools)
    print(f"d_sample_400_*.xlsx: {len(tools)} tools, {missing} without a description")


if __name__ == "__main__":
    main()
